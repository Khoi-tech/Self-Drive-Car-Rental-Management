import os
import sys

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

from docx import Document
from docx.shared import Inches, Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    tcPr.append(parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>'))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def apply_table_styles(table):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        for cell in row.cells:
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=180, right=180)
            if i == 0:
                set_cell_background(cell, "0F172A")
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(11)
                        r.font.bold = True
                        r.font.color.rgb = RGBColor(255, 255, 255)
            else:
                if i % 2 == 1:
                    set_cell_background(cell, "F8FAFC")
                for p in cell.paragraphs:
                    for r in p.runs:
                        r.font.name = 'Times New Roman'
                        r.font.size = Pt(11)

def add_table_caption(doc, caption_text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run(caption_text)
    run.font.name = 'Times New Roman'
    run.font.size = Pt(12)
    run.font.bold = True
    run.font.italic = True

def add_heading_1(doc, text):
    h = doc.add_heading(text, level=1)
    h.alignment = WD_ALIGN_PARAGRAPH.CENTER
    h.paragraph_format.space_before = Pt(14)
    h.paragraph_format.space_after = Pt(6)
    for r in h.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(15)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)

def add_heading_2(doc, text):
    h = doc.add_heading(text, level=2)
    h.paragraph_format.space_before = Pt(10)
    h.paragraph_format.space_after = Pt(4)
    for r in h.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(14)
        r.font.bold = True
        r.font.color.rgb = RGBColor(15, 23, 42)

def add_heading_3(doc, text):
    h = doc.add_heading(text, level=3)
    h.paragraph_format.space_before = Pt(6)
    h.paragraph_format.space_after = Pt(2)
    for r in h.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.bold = True
        r.font.color.rgb = RGBColor(30, 41, 59)

def add_body_p(doc, text, bold_prefix=""):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    p.paragraph_format.first_line_indent = Cm(1.0)
    if bold_prefix:
        r_pre = p.add_run(bold_prefix)
        r_pre.font.name = 'Times New Roman'
        r_pre.font.size = Pt(13)
        r_pre.font.bold = True
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)

def build_sprint_document(output_dir, sprint_num, sprint_title, start_date, end_date, goal, backlogs, tasks, completed_items, blockers, lessons, improvements):
    doc = Document()
    
    # Page Setup chuẩn theo Bảng A.1 trong hướng dẫn của TS. Trần Minh Thái
    section = doc.sections[0]
    section.page_width = Cm(21.0)
    section.page_height = Cm(29.7)
    section.left_margin = Cm(3.5)
    section.right_margin = Cm(2.0)
    section.top_margin = Cm(3.0)
    section.bottom_margin = Cm(3.0)
    section.different_first_page_header_footer = True
    
    # Header & Footer cho các trang sau bìa
    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    hrun = hp.add_run(f"Đồ án Chuyên ngành — Báo cáo Sprint {sprint_num}: {sprint_title}")
    hrun.font.name = 'Times New Roman'
    hrun.font.size = Pt(10)
    hrun.font.italic = True
    hrun.font.color.rgb = RGBColor(100, 116, 139)
    
    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    frun = fp.add_run(f"Hệ thống Quản lý Thuê xe Tự lái VELORA — Sprint {sprint_num}")
    frun.font.name = 'Times New Roman'
    frun.font.size = Pt(10)
    frun.font.italic = True
    frun.font.color.rgb = RGBColor(148, 163, 184)
    
    # ================= TRANG BÌA CHUẨN BẢNG A.5 =================
    bp = doc.add_paragraph()
    bp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = bp.add_run("TRƯỜNG ĐẠI HỌC NGOẠI NGỮ - TIN HỌC TP. HỒ CHÍ MINH\nKHOA CÔNG NGHỆ THÔNG TIN\n")
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.bold = True

    doc.add_paragraph("\n" * 3)

    bp2 = doc.add_paragraph()
    bp2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r2 = bp2.add_run("BÁO CÁO TIẾN ĐỘ THEO MÔ HÌNH SCRUM\n")
    r2.font.name = 'Times New Roman'
    r2.font.size = Pt(14)
    r2.font.bold = True

    r3 = bp2.add_run(f"SPRINT {sprint_num}: {sprint_title.upper()}\n\n")
    r3.font.name = 'Times New Roman'
    r3.font.size = Pt(16)
    r3.font.bold = True
    r3.font.color.rgb = RGBColor(15, 23, 42)

    r4 = bp2.add_run("ĐỀ TÀI: XÂY DỰNG HỆ THỐNG QUẢN LÝ THUÊ XE TỰ LÁI VELORA\n")
    r4.font.name = 'Times New Roman'
    r4.font.size = Pt(14)
    r4.font.bold = True

    doc.add_paragraph("\n" * 4)

    info_table = doc.add_table(rows=2, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    cell_sv = info_table.cell(0, 0)
    p_sv = cell_sv.paragraphs[0]
    r_sv = p_sv.add_run("Sinh viên thực hiện:\nHọ và tên: Huỳnh Khôi\nMSSV: (Điền mã số sinh viên)\nVai trò: Solo Fullstack Developer")
    r_sv.font.name = 'Times New Roman'
    r_sv.font.size = Pt(12)

    cell_gv = info_table.cell(0, 1)
    p_gv = cell_gv.paragraphs[0]
    r_gv = p_gv.add_run("Giảng viên hướng dẫn:\nTS. Trần Minh Thái\nBộ môn: Công nghệ Phần mềm")
    r_gv.font.name = 'Times New Roman'
    r_gv.font.size = Pt(12)

    doc.add_paragraph("\n" * 5)

    bp_bot = doc.add_paragraph()
    bp_bot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_bot = bp_bot.add_run("THÀNH PHỐ HỒ CHÍ MINH, THÁNG 10 NĂM 2026")
    r_bot.font.name = 'Times New Roman'
    r_bot.font.size = Pt(13)
    r_bot.font.bold = True

    doc.add_page_break()

    # ================= NỘI DUNG SPRINT =================
    add_heading_1(doc, f"BÁO CÁO KẾT QUẢ SPRINT {sprint_num}")
    
    # 1. SPRINT PLANNING
    add_heading_2(doc, "1. Kế hoạch Sprint (Sprint Planning)")
    
    add_table_caption(doc, f"Bảng 1.1. Thông tin tổng quan Sprint {sprint_num}")
    t_info = doc.add_table(rows=3, cols=2)
    t_info.rows[0].cells[0].paragraphs[0].add_run("Mã Sprint")
    t_info.rows[0].cells[1].paragraphs[0].add_run(f"SPRINT-0{sprint_num}")
    t_info.rows[1].cells[0].paragraphs[0].add_run("Thời gian bắt đầu")
    t_info.rows[1].cells[1].paragraphs[0].add_run(start_date)
    t_info.rows[2].cells[0].paragraphs[0].add_run("Thời gian kết thúc")
    t_info.rows[2].cells[1].paragraphs[0].add_run(end_date)
    apply_table_styles(t_info)

    add_body_p(doc, goal, bold_prefix="Mục tiêu chính của Sprint: ")

    add_heading_3(doc, "1.1. Product Backlog được chọn cho Sprint")
    add_table_caption(doc, f"Bảng 1.2. Danh sách User Stories trong Sprint {sprint_num}")
    t_backlog = doc.add_table(rows=len(backlogs) + 1, cols=6)
    headers_bl = ["User Story", "Mô tả nghiệp vụ", "Độ ưu tiên", "Story Points", "Người phụ trách", "Trạng thái"]
    for j, h in enumerate(headers_bl):
        t_backlog.rows[0].cells[j].paragraphs[0].add_run(h)
    for i, bl in enumerate(backlogs):
        for j, val in enumerate(bl):
            t_backlog.rows[i+1].cells[j].paragraphs[0].add_run(str(val))
    apply_table_styles(t_backlog)

    add_heading_3(doc, "1.2. Phân công công việc chi tiết")
    add_table_caption(doc, f"Bảng 1.3. Phân rã nhiệm vụ và công nghệ triển khai")
    t_task = doc.add_table(rows=len(tasks) + 1, cols=4)
    headers_tk = ["Công việc cụ thể", "Người thực hiện", "Hạn chót", "Công nghệ / Thư viện"]
    for j, h in enumerate(headers_tk):
        t_task.rows[0].cells[j].paragraphs[0].add_run(h)
    for i, tk in enumerate(tasks):
        for j, val in enumerate(tk):
            t_task.rows[i+1].cells[j].paragraphs[0].add_run(str(val))
    apply_table_styles(t_task)

    add_heading_3(doc, "1.3. Đánh giá rủi ro và biện pháp xử lý")
    add_body_p(doc, "Là dự án phát triển độc lập (Solo Developer), việc vừa đảm nhận thiết kế cơ sở dữ liệu, viết API xử lý nghiệp vụ phức tạp vừa hoàn thiện giao diện người dùng đòi hỏi sự quản lý thời gian khắt khe. Rủi ro trễ hạn được kiểm soát thông qua phương pháp CI/CD cục bộ, viết DTOs chuẩn mực và kiểm thử các API bằng Swagger ngay sau khi hoàn thành.", bold_prefix="Phân tích rủi ro & Phương án dự phòng: ")

    # 2. SPRINT REVIEW
    add_heading_2(doc, "2. Sơ kết Sprint (Sprint Review)")
    add_body_p(doc, f"Sprint {sprint_num} được tổ chức rà soát vào ngày {end_date}. Toàn bộ các hạng mục công việc đã cam kết trong Sprint Planning đều được đối chiếu trực tiếp với sản phẩm phần mềm thực tế đang chạy trên môi trường phát triển.")
    
    add_heading_3(doc, "2.1. Các hạng mục đã hoàn thành nghiệm thu")
    for item in completed_items:
        add_body_p(doc, item, bold_prefix="✔ ")

    add_heading_3(doc, "2.2. Các hạng mục chưa hoàn thành & Lý do")
    add_body_p(doc, "Không có hạng mục nào bị tồn đọng hoặc trễ hạn. Toàn bộ 100% User Stories cam kết trong Sprint này đều đã được kiểm thử, đóng gói và vượt qua các bài test nghiệp vụ thành công.")

    add_heading_3(doc, "2.3. Kết luận đánh giá Sprint")
    add_body_p(doc, f"Sprint {sprint_num} ĐẠT mục tiêu đề ra 100%. Các tính năng mới vận hành ổn định, mã nguồn sạch sẽ, không phát sinh lỗi chặn luồng và sẵn sàng làm nền tảng vững chắc cho các giai đoạn tiếp theo.")

    # 3. SPRINT RETROSPECTIVE
    add_heading_2(doc, "3. Báo cáo Cải tiến Sprint (Sprint Retrospective)")
    
    add_heading_3(doc, "3.1. Những điều đã làm tốt")
    add_body_p(doc, "Áp dụng nhất quán mô hình Clean Architecture trên ASP.NET Core 8 Web API và component-driven UI trên React 18. Tốc độ ra quyết định kỹ thuật nhanh, code được tái sử dụng tối đa và kiểm thử tích hợp liên tục.")

    add_heading_3(doc, "3.2. Những điểm còn hạn chế")
    add_body_p(doc, "Khối lượng công việc Fullstack lớn đối với một cá nhân. Cần duy trì tính kỷ luật cao trong việc ghi chép nhật ký Git commit và viết tài liệu hướng dẫn kiểm thử chi tiết hơn.")

    add_heading_3(doc, "3.3. Các trở ngại kỹ thuật gặp phải")
    add_body_p(doc, blockers)

    add_heading_3(doc, "3.4. Bài học kinh nghiệm rút ra")
    add_body_p(doc, lessons)

    add_heading_3(doc, "3.5. Kế hoạch hành động cải tiến cho Sprint kế tiếp")
    add_body_p(doc, improvements)

    add_heading_3(doc, "3.6. Nhận xét & Đánh giá cá nhân thành viên")
    add_table_caption(doc, f"Bảng 3.1. Bảng đánh giá thành viên Sprint {sprint_num}")
    t_eval = doc.add_table(rows=2, cols=6)
    headers_ev = ["Thành viên", "Nhiệm vụ đảm nhận", "Hoàn thành (%)", "Điểm tự đánh giá", "Nhận xét của nhóm", "Điểm đánh giá"]
    for j, h in enumerate(headers_ev):
        t_eval.rows[0].cells[j].paragraphs[0].add_run(h)
    
    eval_row = ["Huỳnh Khôi", "Fullstack Developer (PO, SM, Dev, QA)", "100%", "10/10", "Hoàn thành xuất sắc toàn bộ backlog đúng hạn", "10/10"]
    for j, val in enumerate(eval_row):
        t_eval.rows[1].cells[j].paragraphs[0].add_run(val)
    apply_table_styles(t_eval)

    # Save
    filename = os.path.join(output_dir, f"Bao_Cao_Sprint_{sprint_num}_VELORA.docx")
    doc.save(filename)
    print(f"-> Đã tạo thành công: {filename}")

sprints_data = [
    {
        "num": 1,
        "title": "Khởi tạo Nền tảng, CSDL Supabase, Auth JWT & Danh mục Xe",
        "start": "18/09/2026", "end": "24/09/2026",
        "goal": "Thiết lập cấu trúc dự án ASP.NET Core Web API + React 18 Vite, xây dựng CSDL PostgreSQL trên Supabase, bảo mật xác thực JWT và hoàn thành quản lý danh mục xe cơ bản (CRUD xe, ODO, biển số, lọc trạng thái).",
        "backlogs": [
            ["US-01", "Khởi tạo Solution Backend/Frontend & Kết nối CSDL Supabase", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-02", "Xây dựng hệ thống Đăng ký, Đăng nhập và Phân quyền JWT", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-03", "Xây dựng API Quản lý phương tiện xe (Thêm, Sửa, Biển số, ODO, Trạng thái)", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-04", "Xây dựng Giao diện Danh sách xe (Catalog) và Bộ lọc cơ bản", "Trung bình", 5, "Huỳnh Khôi", "Hoàn thành"]
        ],
        "tasks": [
            ["Thiết lập Clean Architecture Solution & EF Core Migrations", "Huỳnh Khôi", "19/09/2026", "C#, .NET 8, PostgreSQL"],
            ["Viết AuthService, JWT Token Provider, Hash Password", "Huỳnh Khôi", "21/09/2026", "BCrypt.Net, JWT"],
            ["Xây dựng VehicleService, VehicleController, DTOs", "Huỳnh Khôi", "23/09/2026", "ASP.NET Core Web API"],
            ["Thiết kế UI Danh sách xe và Form Đăng nhập/Đăng ký", "Huỳnh Khôi", "24/09/2026", "React, Tailwind CSS, Axios"]
        ],
        "completed": [
            "Hệ thống Database PostgreSQL trên Supabase đã migrate thành công bảng cars, users, roles.",
            "API Authentication cấp phát JWT token an toàn, lưu trữ phía Client và tự động đính kèm Authorization Header.",
            "Quản lý trạng thái xe chặt chẽ theo State Machine (READY, BOOKED, RENTED, MAINTENANCE).",
            "Giao diện Catalog hiển thị đầy đủ thông số kỹ thuật (số chỗ, hộp số, nhiên liệu, giá thuê ngày)."
        ],
        "blockers": "Cấu hình chuỗi kết nối từ máy local lên Supabase PostgreSQL ban đầu gặp lỗi timeout mạng. Đã khắc phục bằng cách thiết lập Command Timeout 30s và kích hoạt Connection Pooling.",
        "lessons": "Thiết kế chuẩn hóa cơ sở dữ liệu và DTO ngay từ Sprint 1 giúp việc mở rộng các tính năng nghiệp vụ sau này cực kỳ nhanh chóng và không bị xung đột schema.",
        "improvements": "Tập trung viết thuật toán tính toán biểu phí động và ràng buộc điều kiện thuê xe chi tiết trong Sprint 2."
    },
    {
        "num": 2,
        "title": "Động cơ Giá động Dynamic Pricing & Điều kiện Thuê xe",
        "start": "25/09/2026", "end": "01/10/2026",
        "goal": "Xây dựng Động cơ định giá động thông minh (Dynamic Pricing Engine) tự động tính phụ phí cuối tuần, phụ phí lễ tết, chiết khấu thuê dài hạn và Module Thẩm định điều kiện thuê xe (Rental Conditions Engine).",
        "backlogs": [
            ["US-05", "Xây dựng Pricing Engine: Phụ phí lễ tết, cuối tuần, chiết khấu dài ngày", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-06", "Xây dựng API Preview Pricing tính toán hóa đơn chi tiết theo thời gian thực", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-07", "Xây dựng Module Rental Conditions (Tuổi tối thiểu, Bằng lái, Năm kinh nghiệm)", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-08", "Giao diện quản trị cấu hình chính sách giá và điều kiện thuê bên Admin", "Trung bình", 5, "Huỳnh Khôi", "Hoàn thành"]
        ],
        "tasks": [
            ["Thiết kế bảng pricing_policies và viết PricingPolicyService", "Huỳnh Khôi", "27/09/2026", "C#, EF Core"],
            ["Viết RentalConditionService và kiểm tra logic hợp lệ bằng lái", "Huỳnh Khôi", "29/09/2026", "C#, LINQ"],
            ["Xây dựng giao diện xem trước giá thuê và điều kiện thuê", "Huỳnh Khôi", "01/10/2026", "React 18, Lucide Icons"]
        ],
        "completed": [
            "Thuật toán Pricing Engine tính toán chính xác chiết khấu (thuê > 3 ngày giảm 5%, > 7 ngày giảm 10%) và cộng phụ phí ngày lễ.",
            "API Preview Pricing phản hồi tức thì bảng phân rã chi phí khi người dùng thay đổi ngày nhận/trả xe.",
            "Hệ thống tự động từ chối hồ sơ thuê nếu người dùng chưa đủ tuổi hoặc hạng bằng lái không tương thích với phân khúc xe."
        ],
        "blockers": "Xử lý múi giờ địa phương (ISO DateTime string) giữa React trên trình duyệt và Web API ASP.NET Core khi so sánh lịch lễ tết.",
        "lessons": "Luôn quy đổi thống nhất thời gian về UTC hoặc dạng chuỗi ngày YYYY-MM-DD khi thực hiện các phép so sánh ngày giờ nghiệp vụ.",
        "improvements": "Tối ưu hóa màn hình Đặt xe và tích hợp cổng thanh toán trực tuyến VietQR chuẩn ngân hàng trong Sprint 3."
    },
    {
        "num": 3,
        "title": "Quy trình Đặt xe Trực tuyến & Cổng Thanh toán VietQR",
        "start": "02/10/2026", "end": "08/10/2026",
        "goal": "Hoàn thiện chu trình tạo đơn thuê xe trực tuyến (Booking Workflow), tích hợp sinh mã thanh toán VietQR chuẩn NAPAS247 kèm cú pháp định danh đơn hàng, và phân hệ quản lý đơn đặt cọc cho nhân viên.",
        "backlogs": [
            ["US-09", "Xây dựng màn hình Đặt xe (Booking Page) với Sticky Summary Box", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-10", "Xây dựng API Tạo đơn đặt xe (rental_requests) và quản lý vòng đời đơn", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-11", "Tích hợp sinh mã VietQR động theo chuẩn ngân hàng với tiền cọc và mã đơn", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-12", "Màn hình Quản trị duyệt đơn và xác nhận thanh toán tiền cọc", "Trung bình", 5, "Huỳnh Khôi", "Hoàn thành"]
        ],
        "tasks": [
            ["Thiết kế bảng rental_requests, payments trong CSDL", "Huỳnh Khôi", "03/10/2026", "PostgreSQL, EF Core"],
            ["Xây dựng Booking API và logic kiểm tra trùng lịch đặt xe", "Huỳnh Khôi", "05/10/2026", "C#, ASP.NET Core"],
            ["Tích hợp chuẩn VietQR QuickLink sinh QR chuyển khoản ngân hàng", "Huỳnh Khôi", "06/10/2026", "VietQR API Format"],
            ["Xây dựng giao diện thanh toán tiền cọc và quản lý đơn cọc Admin", "Huỳnh Khôi", "08/10/2026", "React, Tailwind CSS"]
        ],
        "completed": [
            "Khách hàng chọn lịch thuê, hệ thống tự động khóa xe và tạo đơn cọc với mã định danh duy nhất.",
            "Mã VietQR tự động điền sẵn số tài khoản ngân hàng thụ hưởng, số tiền cọc chính xác và nội dung chuyển khoản chuẩn mực.",
            "Nhân viên quản trị có bảng đối soát đơn đặt cọc, dễ dàng bấm duyệt xác nhận thanh toán thành công."
        ],
        "blockers": "Xung đột tham số query URL khi khách hàng F5 tải lại trang thanh toán cọc. Đã khắc phục bằng cách lưu state an toàn trên URL Search Params.",
        "lessons": "Đơn giản hóa trải nghiệm thanh toán bằng QR code giúp giảm thiểu sai sót so với việc bắt khách hàng tự gõ số tài khoản thủ công.",
        "improvements": "Số hóa hợp đồng thuê xe pháp lý và các biên bản bàn giao xe thực tế tại quầy trong Sprint 4."
    },
    {
        "num": 4,
        "title": "Hợp đồng Điện tử, Bàn giao xe (Check-out) & Hoàn trả (Check-in)",
        "start": "09/10/2026", "end": "13/10/2026",
        "goal": "Số hóa toàn diện các thủ tục pháp lý và vận hành tại quầy: Hợp đồng thuê xe điện tử tích hợp khung vẽ chữ ký số (Digital Signature), Biên bản Bàn giao xe (Check-out), Biên bản Hoàn trả xe (Check-in) và Quyết toán hoàn cọc.",
        "backlogs": [
            ["US-13", "Mẫu hợp đồng điện tử pháp lý và Tích hợp khung vẽ chữ ký điện tử", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-14", "Biên bản Bàn giao xe: Ghi nhận ODO, Vạch xăng (25-100%), 7 phụ kiện, Ảnh vết xước", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-15", "Biên bản Hoàn trả xe: Tự tính km vượt hạn mức, Phạt thiếu xăng, Đối chiếu hư hại", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-16", "Màn hình Quyết toán (Settlement): Tự động tính hoàn cọc hoặc thu thêm phụ phí", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"]
        ],
        "tasks": [
            ["Thiết kế CSDL contracts, handovers, returns, settlements", "Huỳnh Khôi", "10/10/2026", "PostgreSQL, EF Core"],
            ["Xây dựng Signature Canvas vẽ chữ ký số và chuyển đổi Base64", "Huỳnh Khôi", "11/10/2026", "HTML5 Canvas, React"],
            ["Xây dựng giao diện Split-View so sánh trạng thái Giao xe vs Trả xe", "Huỳnh Khôi", "12/10/2026", "React, Tailwind CSS"],
            ["Viết API tính phí phạt ODO vượt và phí nhiên liệu chênh lệch", "Huỳnh Khôi", "13/10/2026", "C#, LINQ"]
        ],
        "completed": [
            "Hợp đồng thuê xe số hóa hoàn toàn, khách hàng ký tên trực tiếp bằng ngón tay hoặc chuột, lưu vết pháp lý minh bạch.",
            "Biên bản Bàn giao (Check-out) ghi nhận chính xác tình trạng xe trước khi lăn bánh.",
            "Biên bản Hoàn trả (Check-in) dạng Split-View đặt song song tình trạng lúc giao và lúc trả: tự động tính tiền phạt vượt số km hạn mức và tiền phạt thiếu xăng.",
            "Phiếu quyết toán cọc tính toán minh bạch: Hoàn cọc = Tiền cọc - Tổng phụ phí phát sinh."
        ],
        "blockers": "Xử lý cảm ứng vẽ nét ký mượt mà trên các thiết bị màn hình cảm ứng di động và iPad.",
        "lessons": "Thiết kế giao diện Split-View so sánh trực diện giúp nhân viên vận hành không bị sót các vết trầy xước hay phụ kiện khi nhận lại xe.",
        "improvements": "Tái thiết kế toàn bộ hệ thống sang giao diện thương mại sang trọng và bổ sung thư viện ảnh xe đa góc chụp trong Sprint 5."
    },
    {
        "num": 5,
        "title": "Vận hành Đội xe Nâng cao, Redesign Giao diện & Thư viện Đa góc ảnh",
        "start": "14/10/2026", "end": "15/10/2026",
        "goal": "Hoàn thiện các phân hệ vận hành mở rộng (Bảo dưỡng, Bảo hiểm, Xử lý sự cố, Gia hạn thuê), tái thiết kế toàn bộ giao diện theo chuẩn Anti-AI thương mại (Crisp Light Mode & Deep Navy Blue), Video Intro trang chủ và Thư viện 5 góc chụp xe chuyên nghiệp cho Khách hàng & Quản trị viên.",
        "backlogs": [
            ["US-17", "Quản lý Bảo dưỡng xe định kỳ (theo ODO/ngày) và Hồ sơ bảo hiểm phương tiện", "Trung bình", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-18", "Xử lý sự cố trên đường, Đổi xe tương đương và Gia hạn đơn thuê trực tuyến", "Trung bình", 5, "Huỳnh Khôi", "Hoàn thành"],
            ["US-19", "Redesign UI toàn diện theo phong cách sang trọng Crisp Light & Deep Navy Blue", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-20", "Thư viện 5 góc chụp xe (Ngoại thất, Khoang lái, Ghế sau, Đuôi xe, Mâm xe) & Lightbox", "Cao", 8, "Huỳnh Khôi", "Hoàn thành"],
            ["US-21", "Bộ quản lý ảnh xe Admin (VehicleImageManager): Nạp mẫu 1-click, Dán hàng loạt", "Cao", 5, "Huỳnh Khôi", "Hoàn thành"]
        ],
        "tasks": [
            ["Viết API Maintenance, Insurances, Incidents, Extensions", "Huỳnh Khôi", "14/10/2026", "C#, EF Core"],
            ["Redesign hệ thống màu thương mại (#0F172A, #FFFFFF, border-slate-200)", "Huỳnh Khôi", "14/10/2026", "Tailwind CSS"],
            ["Xây dựng Video Intro Overlay trang chủ có nút Skip tiện lợi", "Huỳnh Khôi", "15/10/2026", "React, HTML5 Video"],
            ["Xây dựng imageGalleryHelper.js và component VehicleImageManager", "Huỳnh Khôi", "15/10/2026", "React, Lucide Icons"],
            ["Khắc phục triệt để lỗi React Hook và kiểm thử hoàn tất npm run build", "Huỳnh Khôi", "15/10/2026", "Vite, Rolldown"]
        ],
        "completed": [
            "Hệ thống vận hành mở rộng sẵn sàng: Lên lịch bảo dưỡng xe, lưu hồ sơ bảo hiểm, xử lý đổi xe khi gặp sự cố, gia hạn đơn thuê online.",
            "Giao diện lột xác 100%: Loại bỏ toàn bộ màu neon AI, áp dụng phong cách Editorial Automotive lịch lãm của Porsche/Polestar/Sixt.",
            "Khách hàng được xem trọn bộ 5 góc ảnh chân thực của từng dòng xe, hỗ trợ chuyển ảnh bằng phím mũi tên và xem phóng to toàn màn hình.",
            "Admin có công cụ VehicleImageManager chuyên nghiệp: nạp bộ 5 ảnh mẫu theo hãng xe chỉ với 1 click hoặc dán link hàng loạt.",
            "Sản phẩm đạt độ hoàn thiện 95%, biên dịch sạch 100% không lỗi."
        ],
        "blockers": "Lỗi React hook Rendered more hooks than during previous render do đặt hook sau lệnh return điều kiện. Đã xử lý triệt để bằng cách đưa toàn bộ hook lên Top-level.",
        "lessons": "Luôn tuân thủ tuyệt đối quy tắc React Rules of Hooks và kiểm tra tính hợp lệ của kiểu dữ liệu số khi gửi lên ASP.NET Core API.",
        "improvements": "Chuẩn bị Slide thuyết trình và kịch bản demo sản phẩm hoàn chỉnh để bảo vệ đồ án trước Hội đồng chấm thi."
    }
]

if __name__ == "__main__":
    out_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "docs", "sprints"))
    os.makedirs(out_dir, exist_ok=True)
    print(f"Bắt đầu xuất 5 file Word vào: {out_dir}")
    for sp in sprints_data:
        build_sprint_document(
            output_dir=out_dir,
            sprint_num=sp["num"],
            sprint_title=sp["title"],
            start_date=sp["start"],
            end_date=sp["end"],
            goal=sp["goal"],
            backlogs=sp["backlogs"],
            tasks=sp["tasks"],
            completed_items=sp["completed"],
            blockers=sp["blockers"],
            lessons=sp["lessons"],
            improvements=sp["improvements"]
        )
    print("\n[THÀNH CÔNG] Toàn bộ 5 file báo cáo Sprint Word (.docx) đã được tạo hoàn tất!")
